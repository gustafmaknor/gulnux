"""Gloffice kärna: läser och ändrar Word-, Excel- och PowerPoint-filer.

Delas av webbgränssnittet (web.py) och MCP-servern (mcp.py). Varje ändring sparar
först en säkerhetskopia, så att den kan ångras med undo().
"""

import hashlib
import os
import re
import shutil
import subprocess
import tempfile
import time
from datetime import date, datetime, time as dtime
from pathlib import Path

import docx
import openpyxl
import pptx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries
from pptx.util import Inches

KINDS = {".docx": "document", ".xlsx": "spreadsheet", ".xlsm": "spreadsheet", ".pptx": "presentation"}
# Äldre format och OpenDocument öppnas genom att konverteras med LibreOffice
CONVERTIBLE = {".doc", ".odt", ".rtf", ".xls", ".ods", ".ppt", ".odp"}
LABELS = {"document": "ett textdokument", "spreadsheet": "ett kalkylark", "presentation": "en presentation"}

HOME = Path.home().resolve()
DOCS_DIR = Path(os.environ.get("GLOFFICE_DIR", HOME / "Document")).resolve()
DATA_DIR = Path(os.environ.get("XDG_DATA_HOME", HOME / ".local" / "share")) / "gloffice"
BACKUPS_KEPT = 20


class GlofficeError(Exception):
    pass


# ---------------------------------------------------------------- filer

def resolve(path):
    """Gör en sökväg absolut (relativa utgår från ~/Document) och håll den inom hemkatalogen."""
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = DOCS_DIR / p
    p = p.resolve()
    if not p.is_relative_to(HOME):
        raise GlofficeError(f"{p} ligger utanför hemkatalogen")
    return p


def kind(p):
    try:
        return KINDS[p.suffix.lower()]
    except KeyError:
        hint = " – konvertera den först med convert" if p.suffix.lower() in CONVERTIBLE else ""
        raise GlofficeError(f"{p.name}: filtypen stöds inte{hint}") from None


def _existing(path, expected=None):
    p = resolve(path)
    if not p.is_file():
        raise GlofficeError(f"{p} finns inte")
    k = kind(p)
    if expected and k != expected:
        raise GlofficeError(f"{p.name} är inte {LABELS[expected]}")
    return p, k


def mtime(path):
    # Som sträng: nanosekunder får inte plats exakt i ett JavaScript-tal
    return str(resolve(path).stat().st_mtime_ns)


def list_documents(directory=None):
    base = resolve(directory) if directory else DOCS_DIR
    base.mkdir(parents=True, exist_ok=True)
    files = []
    for f in sorted(base.rglob("*")):
        rel = f.relative_to(base)
        suffix = f.suffix.lower()
        if any(part.startswith((".", "~$")) for part in rel.parts) or not f.is_file():
            continue
        if suffix in KINDS or suffix in CONVERTIBLE:
            files.append({
                "path": str(f),
                "name": rel.as_posix(),
                "kind": KINDS.get(suffix, "convertible"),
                "modified": f.stat().st_mtime,
            })
            if len(files) >= 500:
                break
    return files


def create(path):
    p = resolve(path)
    k = kind(p)
    if p.exists():
        raise GlofficeError(f"{p} finns redan")
    p.parent.mkdir(parents=True, exist_ok=True)
    if k == "document":
        docx.Document().save(p)
    elif k == "spreadsheet":
        openpyxl.Workbook().save(p)
    else:
        prs = pptx.Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
        prs.save(p)
    return str(p)


# ---------------------------------------------------------------- ändringar och ångra

def _backup_dir(p):
    # Kort namn: filnamnet plus en hash av hela sökvägen, så att lika namn i olika mappar hålls isär
    digest = hashlib.sha1(str(p).encode()).hexdigest()[:10]
    return DATA_DIR / "backups" / f"{p.name}-{digest}"


def _backup(p):
    d = _backup_dir(p)
    d.mkdir(parents=True, exist_ok=True)
    shutil.copy2(p, d / f"{time.time_ns()}{p.suffix}")
    for old in sorted(d.iterdir())[:-BACKUPS_KEPT]:
        old.unlink()


def undo(path):
    p, _ = _existing(path)
    d = _backup_dir(p)
    backups = sorted(d.iterdir()) if d.is_dir() else []
    if not backups:
        raise GlofficeError(f"Det finns inget att ångra i {p.name}")
    shutil.copyfile(backups[-1], p)
    backups[-1].unlink()
    return f"Ångrade senaste ändringen i {p.name}"


def _load(p, k):
    if k == "document":
        return docx.Document(p)
    if k == "spreadsheet":
        return openpyxl.load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
    return pptx.Presentation(p)


def _modify(path, expected, change):
    """Öppna filen, kör change(objekt), spara en säkerhetskopia och skriv tillbaka."""
    p, k = _existing(path, expected)
    obj = _load(p, k)
    result = change(obj)
    _backup(p)
    obj.save(p)
    return result


# ---------------------------------------------------------------- läsa

def read(path):
    """Strukturerat innehåll för webbgränssnittet."""
    p, k = _existing(path)
    data = {"path": str(p), "name": p.name, "kind": k, "mtime": mtime(p)}
    data.update({"document": _read_document, "spreadsheet": _read_spreadsheet,
                 "presentation": _read_presentation}[k](p))
    return data


def _blocks(doc):
    """Stycken och tabeller i dokumentordning. Styckenas index matchar doc.paragraphs."""
    pi = ti = 0
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            par = Paragraph(el, doc._body)
            style = par.style.name if par.style is not None else ""
            yield {"type": "paragraph", "index": pi, "style": style, "text": par.text}
            pi += 1
        elif el.tag == qn("w:tbl"):
            table = Table(el, doc._body)
            yield {"type": "table", "index": ti, "rows": [[c.text for c in r.cells] for r in table.rows]}
            ti += 1


def _read_document(p):
    return {"blocks": list(_blocks(docx.Document(p)))}


def _jsonable(v):
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def _formula(v):
    text = getattr(v, "text", v)  # ArrayFormula har formeln i .text
    return text if isinstance(text, str) and text.startswith("=") else None


def _workbooks(p):
    # Två vyer: formlerna som de står, och de senast beräknade värdena
    return openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)


def _read_spreadsheet(p, max_rows=200, max_cols=50):
    wb, values = _workbooks(p)
    sheets = []
    for ws in wb.worksheets:
        vs = values[ws.title]
        rows = [
            [{"v": _jsonable(vs.cell(r, c).value), "f": _formula(ws.cell(r, c).value)}
             for c in range(1, min(ws.max_column, max_cols) + 1)]
            for r in range(1, min(ws.max_row, max_rows) + 1)
        ]
        sheets.append({"name": ws.title, "rows": rows, "max_row": ws.max_row, "max_col": ws.max_column})
    return {"sheets": sheets}


def _read_presentation(p):
    prs = pptx.Presentation(p)
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title
        shapes = [
            {"index": i, "name": shape.name, "text": shape.text_frame.text,
             "title": title is not None and shape.shape_id == title.shape_id}
            for i, shape in enumerate(slide.shapes) if shape.has_text_frame
        ]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        slides.append({"number": number, "layout": slide.slide_layout.name, "shapes": shapes, "notes": notes})
    return {"slides": slides, "layouts": [layout.name for layout in prs.slide_layouts]}


def _cell_summary(c):
    if c["f"]:
        return c["f"] if c["v"] is None else f"{c['v']} ({c['f']})"
    return "" if c["v"] is None else str(c["v"])


def summary(path):
    """Läsbar översikt för agenter, med de index som verktygen använder."""
    d = read(path)
    out = [f"# {d['name']} ({d['path']})"]
    if d["kind"] == "document":
        for b in d["blocks"]:
            if b["type"] == "paragraph":
                out.append(f"[{b['index']}] ({b['style']}) {b['text']}")
            else:
                out.append(f"[tabell {b['index']}]")
                out.extend("| " + " | ".join(row) + " |" for row in b["rows"])
    elif d["kind"] == "spreadsheet":
        for s in d["sheets"]:
            out.append(f"\n## Blad \"{s['name']}\" ({s['max_row']} rader × {s['max_col']} kolumner)")
            for r, row in enumerate(s["rows"][:30], start=1):
                cells = [_cell_summary(c) for c in row]
                out.append(f"{r}\t" + "\t".join(cells))
            if s["max_row"] > 30:
                out.append("… använd sheet_read för resten")
    else:
        out.append(f"Layouter: {', '.join(f'[{i}] {n}' for i, n in enumerate(d['layouts']))}")
        for s in d["slides"]:
            out.append(f"\n## Bild {s['number']} ({s['layout']})")
            for sh in s["shapes"]:
                label = "Rubrik" if sh["title"] else sh["name"]
                out.append(f"  [form {sh['index']}] {label}: {sh['text']}")
            if s["notes"]:
                out.append(f"  Anteckningar: {s['notes']}")
    return "\n".join(out)


# ---------------------------------------------------------------- textdokument

def _paragraph(doc, index):
    pars = doc.paragraphs
    if not 0 <= index < len(pars):
        raise GlofficeError(f"Stycke {index} finns inte (dokumentet har {len(pars)} stycken)")
    return pars[index]


def _set_text(par, text):
    """Byt text men behåll formateringen från första textavsnittet."""
    for link in par._p.findall(qn("w:hyperlink")):
        par._p.remove(link)
    runs = par.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        par.add_run(text)


def _set_style(par, style):
    try:
        par.style = style
    except KeyError:
        raise GlofficeError(f"Formatmallen \"{style}\" finns inte i dokumentet") from None


def docx_replace_paragraph(path, index, text, style=None):
    def change(doc):
        par = _paragraph(doc, index)
        _set_text(par, text)
        if style:
            _set_style(par, style)
        return f"Ändrade stycke {index}"
    return _modify(path, "document", change)


def docx_insert_paragraph(path, text, style=None, after=None):
    def change(doc):
        if after is None:
            par = doc.add_paragraph(text)
            where = len(doc.paragraphs) - 1
        elif after < 0:
            par = _paragraph(doc, 0).insert_paragraph_before(text)
            where = 0
        else:
            anchor = _paragraph(doc, after)
            new = OxmlElement("w:p")
            anchor._p.addnext(new)
            par = Paragraph(new, anchor._parent)
            par.add_run(text)
            where = after + 1
        if style:
            _set_style(par, style)
        return f"La till stycke {where}"
    return _modify(path, "document", change)


def docx_delete_paragraph(path, index):
    def change(doc):
        p = _paragraph(doc, index)._p
        p.getparent().remove(p)
        return f"Tog bort stycke {index}"
    return _modify(path, "document", change)


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def docx_find_replace(path, find, replace):
    if not find:
        raise GlofficeError("Söktexten får inte vara tom")

    def change(doc):
        count = 0
        for par in _all_paragraphs(doc):
            n = par.text.count(find)
            if not n:
                continue
            done = 0
            for run in par.runs:
                done += run.text.count(find)
                run.text = run.text.replace(find, replace)
            if done < n:  # träffen är uppdelad på flera textavsnitt
                _set_text(par, par.text.replace(find, replace))
            count += n
        return f"Ersatte {count} förekomster"
    return _modify(path, "document", change)


def docx_add_table(path, rows, after=None, style=None):
    if not rows or not all(isinstance(r, list) for r in rows):
        raise GlofficeError("rows ska vara en lista med rader, t.ex. [[\"Namn\", \"Pris\"], [\"Kaffe\", \"35\"]]")

    def change(doc):
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        try:
            table.style = style or "Table Grid"
        except KeyError:
            pass  # mallen saknas i dokumentet – behåll standard
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = "" if value is None else str(value)
        if after is not None:
            _paragraph(doc, after)._p.addnext(table._tbl)
        return f"La till en tabell med {len(rows)} rader och {cols} kolumner"
    return _modify(path, "document", change)


# ---------------------------------------------------------------- kalkylark

def _sheet(wb, name):
    if name is None:
        return wb.active
    if name not in wb.sheetnames:
        raise GlofficeError(f"Bladet \"{name}\" finns inte (finns: {', '.join(wb.sheetnames)})")
    return wb[name]


def sheet_read(path, sheet=None, cells=None, max_rows=500):
    p, _ = _existing(path, "spreadsheet")
    wb, values = _workbooks(p)
    ws = _sheet(wb, sheet)
    vs = values[ws.title]
    if cells:
        min_col, min_row, max_col, max_row = range_boundaries(cells.upper())
    else:
        min_col = min_row = 1
        max_col = max_row = None
    min_col, min_row = min_col or 1, min_row or 1
    max_col = max_col or ws.max_column
    max_row = min(max_row or ws.max_row, min_row + max_rows - 1)
    values_out, formulas = [], {}
    for r in range(min_row, max_row + 1):
        row = []
        for c in range(min_col, max_col + 1):
            row.append(_jsonable(vs.cell(r, c).value))
            f = _formula(ws.cell(r, c).value)
            if f:
                formulas[f"{get_column_letter(c)}{r}"] = f
        values_out.append(row)
    return {
        "sheet": ws.title,
        "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
        "values": values_out,
        "formulas": formulas,
        "note": "Formler som skrivits av Gloffice saknar beräknat värde tills recalculate körs.",
    }


def parse_input(text):
    """Tolka det användaren skriver i en cell: tal (med , eller .), formel eller text."""
    if text is None or text == "":
        return None
    if not isinstance(text, str) or text.startswith("="):
        return text
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d*[.,]\d+", text):
        return float(text.replace(",", "."))
    return text


def sheet_write(path, start, values, sheet=None):
    if not isinstance(values, list) or not all(isinstance(r, list) for r in values):
        raise GlofficeError("values ska vara en lista med rader, t.ex. [[1, 2], [\"=A1+B1\"]]")
    col_letter, row0 = coordinate_from_string(start.upper())
    col0 = column_index_from_string(col_letter)

    def change(wb):
        ws = _sheet(wb, sheet)
        for r, row in enumerate(values):
            for c, value in enumerate(row):
                ws.cell(row0 + r, col0 + c).value = value
        width = max((len(r) for r in values), default=1)
        end = f"{get_column_letter(col0 + width - 1)}{row0 + len(values) - 1}"
        return f"Skrev till {ws.title}!{start.upper()}:{end}"
    return _modify(path, "spreadsheet", change)


def sheet_add(path, name):
    def change(wb):
        if name in wb.sheetnames:
            raise GlofficeError(f"Bladet \"{name}\" finns redan")
        wb.create_sheet(name)
        return f"La till bladet \"{name}\""
    return _modify(path, "spreadsheet", change)


# ---------------------------------------------------------------- presentationer

def _slide(prs, number):
    if not 1 <= number <= len(prs.slides):
        raise GlofficeError(f"Bild {number} finns inte (presentationen har {len(prs.slides)} bilder)")
    return prs.slides[number - 1]


def _layout(prs, layout):
    layouts = prs.slide_layouts
    if layout is None:
        return layouts[1] if len(layouts) > 1 else layouts[0]  # oftast "Rubrik och innehåll"
    if isinstance(layout, int):
        if not 0 <= layout < len(layouts):
            raise GlofficeError(f"Layout {layout} finns inte")
        return layouts[layout]
    found = layouts.get_by_name(layout)
    if found is None:
        raise GlofficeError(f"Layouten \"{layout}\" finns inte")
    return found


def slide_add(path, title, bullets=None, layout=None):
    def change(prs):
        slide = prs.slides.add_slide(_layout(prs, layout))
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        if bullets:
            body = next((ph for ph in slide.placeholders
                         if ph.placeholder_format.idx != 0 and ph.has_text_frame), None)
            if body is None:
                body = slide.shapes.add_textbox(Inches(0.7), Inches(1.6),
                                                prs.slide_width - Inches(1.4), prs.slide_height - Inches(2.2))
            tf = body.text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
        return f"La till bild {len(prs.slides)}"
    return _modify(path, "presentation", change)


def slide_set_text(path, slide, shape, text):
    def change(prs):
        shapes = _slide(prs, slide).shapes
        if not 0 <= shape < len(shapes) or not shapes[shape].has_text_frame:
            raise GlofficeError(f"Bild {slide} har ingen textform med index {shape}")
        shapes[shape].text_frame.text = text
        return f"Ändrade text på bild {slide}"
    return _modify(path, "presentation", change)


def slide_delete(path, slide):
    def change(prs):
        _slide(prs, slide)
        ids = prs.slides._sldIdLst
        sld = ids[slide - 1]
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
        return f"Tog bort bild {slide}"
    return _modify(path, "presentation", change)


# ---------------------------------------------------------------- LibreOffice

def _soffice(args):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise GlofficeError("LibreOffice saknas – det behövs för konvertering och omräkning")
    profile = DATA_DIR / "libreoffice-profil"
    try:
        subprocess.run([exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", *args],
                       check=True, capture_output=True, timeout=180)
    except subprocess.CalledProcessError as e:
        raise GlofficeError(f"LibreOffice misslyckades: {e.stderr.decode(errors='replace').strip()}") from None
    except subprocess.TimeoutExpired:
        raise GlofficeError("LibreOffice tog för lång tid") from None


def convert(path, format, out_dir=None):
    """Konvertera med LibreOffice, t.ex. till pdf, docx, xlsx, pptx, odt eller csv."""
    p = resolve(path)
    if not p.is_file():
        raise GlofficeError(f"{p} finns inte")
    out_dir = Path(out_dir) if out_dir else p.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    _soffice(["--convert-to", format, "--outdir", str(out_dir), str(p)])
    out = out_dir / f"{p.stem}.{format.split(':')[0]}"
    if not out.is_file():
        raise GlofficeError("Konverteringen gav ingen fil")
    return str(out)


def has_libreoffice():
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def recalculate(path, backup=True):
    """Räkna om alla formler i ett kalkylark (openpyxl kan inte räkna själv)."""
    p, _ = _existing(path, "spreadsheet")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(convert(p, p.suffix[1:], out_dir=tmp))
        if backup:
            _backup(p)
        shutil.copyfile(out, p)
    return f"Räknade om formlerna i {p.name}"
